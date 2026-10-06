"""Bouwt assets/excel/plaatsingen.xlsx: leeg Q4S-sjabloon voor de export 'Plaatsingen'.

env -u PYTHONPATH uv run -q --isolated --with openpyxl --with pillow python scripts/plaatsingen-sjabloon.py
Het dashboard (src/lib/plaatsingen-excel.ts) zet alleen waarden in de voorbereide cellen.
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter as L
from PIL import Image as PILImage

LOGO = "public/logo/q4slogoOriginalNEW.jpg"
UIT = "assets/excel/plaatsingen.xlsx"
EERSTE, REGELS = 6, 300  # eerste datarij, aantal voorbereide regels (= PLAATSINGEN_SJABLOON in de TS)

ZWART, GRIJS, LICHT, LIJN = "1C1C1E", "6B6B70", "F4F4F5", "D9D9DB"
def font(**k): return Font(name="Calibri", **{"size": 10, **k})
fill = lambda c: PatternFill("solid", fgColor=c)
dun = Side(style="thin", color=LIJN)
EURO = '€ #,##0.00;[Red]-€ #,##0.00;"–"'
DATUM = "dd-mm-yyyy"

# (kop, breedte, formaat)
KOL = [
    ("Werknemer", 22, None), ("Status", 13, None), ("Klant", 24, None), ("Functie", 22, None),
    ("Werklocatie", 18, None), ("Start", 11, DATUM), ("Einde", 11, DATUM), ("Per", 6, None),
    ("Inkoop", 11, EURO), ("Verkoop", 11, EURO), ("Marge", 11, EURO), ("All-in", 7, None),
    ("Km inkoop", 10, EURO), ("Km verkoop", 10, EURO),
    ("Bedrijfsnaam", 24, None), ("KvK", 11, None), ("Btw-nummer", 17, None), ("IBAN", 22, None),
    ("E-mail", 30, None), ("Telefoon", 14, None), ("Adres", 24, None), ("Postcode", 9, None),
    ("Plaats", 14, None), ("PO-nummer", 12, None),
]
GROEPEN = [("Plaatsing & tarief", 1, 14), ("Werknemer — factuur- & contactgegevens", 15, 24)]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Plaatsingen"
ws.sheet_view.showGridLines = False

# Titelband met logo
im = PILImage.open(LOGO).convert("RGB")
h = 52
im.thumbnail((int(h * im.width / im.height) * 2, h * 2))
buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
logo = XLImage(buf); logo.height = h; logo.width = int(h * im.width / im.height)
ws.add_image(logo, "A1")
ws.row_dimensions[1].height = 22; ws.row_dimensions[2].height = 22; ws.row_dimensions[3].height = 10
ws.cell(1, 3, "Plaatsingen").font = font(size=18, bold=True, color=ZWART)
ws.cell(2, 3, "Lopende plaatsingen").font = font(size=10, color=GRIJS)
for c in range(1, len(KOL) + 1):
    ws.cell(3, c).border = Border(bottom=Side(style="medium", color=ZWART))

# Groepskop (rij 4) en kolomkoppen (rij 5)
for naam, van, tot in GROEPEN:
    ws.merge_cells(start_row=4, start_column=van, end_row=4, end_column=tot)
    c = ws.cell(4, van, naam)
    c.font = font(size=9, bold=True, color=GRIJS); c.alignment = Alignment(vertical="center")
ws.row_dimensions[4].height = 20; ws.row_dimensions[5].height = 22
for i, (kop, breedte, fmt) in enumerate(KOL, 1):
    c = ws.cell(5, i, kop)
    c.font = font(bold=True, color="FFFFFF"); c.fill = fill(ZWART)
    c.alignment = Alignment(horizontal="right" if fmt == EURO else "left", vertical="center")
    ws.column_dimensions[L(i)].width = breedte

# Voorbereide regels: opmaak + marge-formule
for r in range(EERSTE, EERSTE + REGELS):
    for i, (kop, _, fmt) in enumerate(KOL, 1):
        c = ws.cell(r, i)
        c.font = font(bold=(i == 1), color=ZWART)
        c.border = Border(bottom=dun)
        if r % 2 == 0: c.fill = fill(LICHT)
        if fmt: c.number_format = fmt
        c.alignment = Alignment(horizontal="right" if fmt == EURO else "left", vertical="center")
    ws.cell(r, 11).value = f'=IF(OR(I{r}="",J{r}=""),"",J{r}-I{r})'
    ws.cell(r, 11).font = font(bold=True, color="0E7A4F")

ws.freeze_panes = "B6"
ws.auto_filter.ref = f"A5:{L(len(KOL))}{EERSTE + REGELS - 1}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "5:5"
wb.save(UIT)
print("ok", UIT)
