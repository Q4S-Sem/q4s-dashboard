"""Bouwt 'Factuur overzicht 2026 (nieuw).xlsx' in Q4S-stijl uit het oude overzicht.

uv run --with openpyxl --with pillow python build.py <oud.xlsx> <nieuw.xlsx> <logo.jpg>
"""
import sys, datetime as dt, re, io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter as L
from PIL import Image as PILImage

OUD, NIEUW, LOGO = sys.argv[1:4]
SJABLOON = OUD == "-"   # leeg sjabloon voor de dashboard-export
EXTRA = 1500 if SJABLOON else 300   # lege, voorbereide regels onder de data
MAXR = 2000          # bereik voor de sommen op Overzicht

# --- Q4S-stijl ---------------------------------------------------------------
ZWART, INKT, GRIJS, LICHT, LIJN = "1C1C1E", "45454B", "6B6B70", "F4F4F5", "D9D9DB"
F = "Calibri"
def font(**k): return Font(name=F, **{"size": 10, **k})
fill = lambda c: PatternFill("solid", fgColor=c)
dun = Side(style="thin", color=LIJN)
rand = Border(bottom=dun)
EURO = '€ #,##0.00;[Red]-€ #,##0.00;"–"'
DATUM = "dd-mm-yyyy"
PCT = "0%"

def logo(ws, hoogte):
    im = PILImage.open(LOGO).convert("RGB")
    im.thumbnail((int(hoogte * im.width / im.height) * 2, hoogte * 2))
    buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
    x = XLImage(buf); x.height = hoogte; x.width = int(hoogte * im.width / im.height)
    ws.add_image(x, "A1")

def kop(ws, titel, sub, breedte, logo_h=52):
    """Titelband: logo links, titel + uitleg ernaast."""
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 22; ws.row_dimensions[2].height = 22; ws.row_dimensions[3].height = 12
    logo(ws, logo_h)
    c = ws.cell(1, 3, titel); c.font = font(size=18, bold=True, color=ZWART)
    c = ws.cell(2, 3, sub); c.font = font(size=10, color=GRIJS)
    for col in range(1, breedte + 1):
        ws.cell(3, col).border = Border(bottom=Side(style="medium", color=ZWART))

# --- oude data lezen ---------------------------------------------------------
wo = None if SJABLOON else openpyxl.load_workbook(OUD, data_only=False).active

def datum(v):
    if isinstance(v, dt.datetime):
        # Een jaartal vóór 2020 is een typfout (bv. 1906): niet gokken, naar de opmerking.
        return (v, None) if v.year >= 2020 else (None, v.strftime("%d-%m-%Y"))
    if isinstance(v, str) and v.strip():
        m = re.fullmatch(r"(\d{1,2})-(\d{1,2})-(\d{4})", v.strip())
        if m:
            try: return dt.datetime(int(m[3]), int(m[2]), int(m[1])), None
            except ValueError: pass
        return None, v.strip()          # onleesbaar → in opmerking, niet verzinnen
    return None, None

def getal(v):
    if isinstance(v, (int, float)): return v
    return None

def tekst(v):
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = str(v).strip()
    return s or None

def btw_bedrag(formule, ex, pct, cache):
    """Formule als het oud ook ex×% was; anders het oude vaste bedrag (bv. btw op eetbon)."""
    if isinstance(formule, str) and formule.startswith("="): return "f"
    if formule is None: return "f"
    if ex is not None and pct is not None and abs(round(ex * pct, 2) - formule) < 0.02: return "f"
    return formule

rijen = []
for r in range(3, 0 if SJABLOON else wo.max_row + 1):
    g = lambda col: wo[f"{col}{r}"].value
    naam = tekst(g("C"))
    if (g("V") or "") == "Totals": continue
    ex_in, ex_uit = getal(g("I")), getal(g("P"))
    verkoop = any(g(c) not in (None, "") for c in "MNOT") or (ex_uit not in (None, 0))
    if not naam and not ex_in and not verkoop: continue      # lege sjabloonregels
    rijen.append(dict(r=r, naam=naam, klant=tekst(g("D")), q4sklant=tekst(g("O")), verkoop=verkoop,
                      fnr=tekst(g("E")), ref=tekst(g("F")), ont=g("G"), bet=g("H"),
                      ex_in=ex_in, pct_in=getal(g("J")), btw_in=g("K"),
                      q4snr=g("M"), ex_uit=ex_uit, pct_uit=getal(g("Q")), btw_uit=g("R"),
                      verst=g("T"), ontv=g("U"), opm=tekst(g("V"))))

freelancers = {x["naam"].lower() for x in rijen if x["verkoop"] and x["naam"]}
facturen = [x for x in rijen if x["verkoop"] or (x["naam"] or "").lower() in freelancers]
kosten = [x for x in rijen if x not in facturen]

SOORTEN = ["Huur & kantoor", "Software & abonnementen", "Auto & reizen", "Verzekeringen",
           "Marketing & werving", "Advies, boekhouder & juridisch", "Bank & rente",
           "Opleiding & certificering", "Belastingen", "Personeel", "Declaraties", "Overig"]
RAAD = [(r"microsoft|snelstart|mbd|kpn|linked|google|adobe|abonnement", "Software & abonnementen"),
        (r"fleetcor|tank|travelcard|lease|shell|parkeer|ns\b|auto", "Auto & reizen"),
        (r"klaverblad|nationale|verzeker|achmea|centraal beheer", "Verzekeringen"),
        (r"belasting", "Belastingen"), (r"bouw|kantoor|huur", "Huur & kantoor"),
        (r"accountant|advies|notaris|juridisch|sna\b|normering|kiwa", "Advies, boekhouder & juridisch"),
        (r"bank|ing\b|rabo|abn", "Bank & rente"), (r"cursus|opleiding|certific|vca", "Opleiding & certificering"),
        (r"spaww|loon|salaris|pensioen", "Personeel"), (r"indeed|vacature|marketing|drukwerk", "Marketing & werving")]
def soort(x):
    t = f"{x['naam'] or ''} {x['klant'] or ''}".lower()
    return next((s for p, s in RAAD if re.search(p, t)), "Overig")

wb = openpyxl.Workbook()

# ============================================================================
# FACTUREN — één regel per factuur van een freelancer, met onze factuur ernaast
# ============================================================================
wf = wb.active; wf.title = "Facturen"
FK = [  # (kop, breedte, groep)
    ("Ontvangen op", 12, "freelancer"), ("Kwartaal", 9, "freelancer"), ("Freelancer", 22, "freelancer"),
    ("Klant / project", 20, "freelancer"), ("Weken / opmerking", 22, "freelancer"),
    ("Factuurnr", 13, "inkoop"), ("Inkoop ex btw", 13, "inkoop"), ("Btw %", 7, "inkoop"),
    ("Btw inkoop", 11, "inkoop"), ("Inkoop incl btw", 14, "inkoop"), ("Betaald op", 12, "inkoop"),
    ("Q4S factuurnr", 13, "verkoop"), ("Verkoop ex btw", 13, "verkoop"), ("Btw % klant", 8, "verkoop"),
    ("Btw verkoop", 11, "verkoop"), ("Verkoop incl btw", 14, "verkoop"), ("Verstuurd op", 12, "verkoop"),
    ("Klant betaald op", 13, "verkoop"),
    ("Marge ex btw", 13, "resultaat"), ("Status", 16, "resultaat"),
]
GROEP = {"freelancer": ("WIE & WAT", INKT), "inkoop": ("INKOOP  ·  factuur van de freelancer", "3F3F46"),
         "verkoop": ("VERKOOP  ·  onze factuur aan de klant", ZWART), "resultaat": ("RESULTAAT", "16A34A")}
kop(wf, "Facturen", "Eén regel per factuur van een freelancer. Vul de witte vakjes in; de grijze rekenen zelf.", len(FK))
GR, HR, D0 = 4, 5, 6
start = 1
for i, (_, b, g) in enumerate(FK, 1):
    wf.column_dimensions[L(i)].width = b
    if i == len(FK) or FK[i][2] != g:
        wf.merge_cells(start_row=GR, start_column=start, end_row=GR, end_column=i)
        c = wf.cell(GR, start, GROEP[g][0]); c.font = font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(start, i + 1): wf.cell(GR, col).fill = fill(GROEP[g][1])
        start = i + 1
wf.row_dimensions[GR].height = 18
for i, (k, _, _) in enumerate(FK, 1):
    c = wf.cell(HR, i, k); c.font = font(bold=True, color=ZWART)
    c.fill = fill(LICHT); c.alignment = Alignment(vertical="center", wrap_text=True)
wf.row_dimensions[HR].height = 30

facturen.sort(key=lambda x: (datum(x["ont"])[0] or datum(x["verst"])[0] or dt.datetime(2100, 1, 1)))
eind_f = D0 + len(facturen) + EXTRA - 1
for n in range(D0, eind_f + 1):
    x = facturen[n - D0] if n - D0 < len(facturen) else None
    if x:
        ont, o1 = datum(x["ont"]); bet, o2 = datum(x["bet"]); verst, o3 = datum(x["verst"]); ontv, o4 = datum(x["ontv"])
        opm = "; ".join(s for s in [x["opm"], x["ref"] and f"ref {x['ref']}",
                                    o1 and f"ontvangen: {o1}", o2 and f"betaald: {o2}",
                                    o3 and f"verstuurd: {o3}", o4 and f"klant betaald: {o4}"] if s)
        klant = x["q4sklant"] or x["klant"]
        if x["q4sklant"] and x["klant"] and x["klant"].lower() not in x["q4sklant"].lower():
            klant = f"{x['q4sklant']} ({x['klant']})"
        vals = {1: ont, 3: x["naam"], 4: klant, 5: opm or None, 6: x["fnr"], 7: x["ex_in"], 8: x["pct_in"],
                11: bet, 12: x["q4snr"], 13: x["ex_uit"], 14: x["pct_uit"], 17: verst, 18: ontv}
        vals[9] = btw_bedrag(x["btw_in"], x["ex_in"], x["pct_in"], None)
        vals[15] = btw_bedrag(x["btw_uit"], x["ex_uit"], x["pct_uit"], None)
    else:
        vals = {9: "f", 15: "f"}
    for col, v in vals.items():
        if v is not None and v != "f": wf.cell(n, col, v)
    f = lambda s: s.replace("#", str(n))
    wf.cell(n, 2, f('=IF(A#="","",YEAR(A#)&" Q"&ROUNDUP(MONTH(A#)/3,0))'))
    if vals.get(9) == "f": wf.cell(n, 9, f('=IF(G#="","",ROUND(G#*H#,2))'))
    wf.cell(n, 10, f('=IF(G#="","",G#+N(I#))'))
    if vals.get(15) == "f": wf.cell(n, 15, f('=IF(M#="","",ROUND(M#*N#,2))'))
    wf.cell(n, 16, f('=IF(M#="","",M#+N(O#))'))
    # Marge per Q4S-factuur: één verkoopfactuur dekt soms meerdere freelancerfacturen;
    # de marge staat dan op de eerste regel van dat factuurnummer.
    wf.cell(n, 19, f('=IF(L#="","",IF(COUNTIF(L$6:L#,L#)>1,"",'
                     'SUMIF(L$6:L$2000,L#,M$6:M$2000)-SUMIF(L$6:L$2000,L#,G$6:G$2000)))'))
    wf.cell(n, 20, f('=IF(AND(C#="",G#=""),"",IF(L#="","Nog factureren",'
                     'IF(AND(G#<>"",OR(K#="",K#>TODAY())),"Nog betalen",'
                     'IF(OR(R#="",R#>TODAY()),"Wacht op klant","Afgerond"))))'))
    for col in range(1, len(FK) + 1):
        c = wf.cell(n, col); c.font = font(color=ZWART); c.border = rand
        c.alignment = Alignment(vertical="center")
        if col in (1, 11, 17, 18): c.number_format = DATUM
        if col in (7, 9, 10, 13, 15, 16, 19): c.number_format = EURO
        if col in (8, 14): c.number_format = PCT
        if col in (2, 9, 10, 15, 16, 19, 20): c.font = font(color=INKT); c.fill = fill("FAFAFA")
    wf.cell(n, 19).font = font(bold=True, color=ZWART)

tab = Table(displayName="tblFacturen", ref=f"A{HR}:{L(len(FK))}{eind_f}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
wf.add_table(tab)
wf.freeze_panes = f"D{D0}"

STATUS = {"Afgerond": ("DCFCE7", "166534"), "Wacht op klant": ("DBEAFE", "1E40AF"),
          "Nog betalen": ("FEF3C7", "92400E"), "Nog factureren": ("FEE2E2", "991B1B")}
def statuskleur(ws, kol, eind):
    for s, (bg, fg) in STATUS.items():
        ws.conditional_formatting.add(f"{kol}{D0}:{kol}{eind}", CellIsRule(
            operator="equal", formula=[f'"{s}"'], fill=fill(bg), font=Font(name=F, size=10, bold=True, color=fg)))
statuskleur(wf, "T", eind_f)

def PCTV():
    v = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    v.prompt, v.promptTitle, v.error = "Typ 21%, 9% of 0%", "Btw-percentage", "Typ een percentage, bv. 21%"
    return v
def DATV():
    v = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    v.error, v.errorTitle = "Typ een datum, bv. 15-10-2026", "Geen datum"
    return v
pct = PCTV(); wf.add_data_validation(pct); pct.add(f"H{D0}:H{eind_f}"); pct.add(f"N{D0}:N{eind_f}")
dat = DATV(); wf.add_data_validation(dat)
for k in "AKQR": dat.add(f"{k}{D0}:{k}{eind_f}")

# ============================================================================
# KOSTEN — alles wat Q4S zelf uitgeeft (geen freelancers)
# ============================================================================
wk = wb.create_sheet("Kosten")
KK = [("Datum", 12), ("Kwartaal", 9), ("Leverancier", 26), ("Omschrijving", 30), ("Soort", 26),
      ("Factuurnr", 16), ("Bedrag ex btw", 13), ("Btw %", 7), ("Btw €", 11), ("Totaal incl btw", 14),
      ("Betaald op", 12), ("Status", 14)]
kop(wk, "Kosten", "Alle eigen kosten van Q4S: huur, software, auto, verzekeringen … Kies bij Soort een categorie.", len(KK))
for i, (k, b) in enumerate(KK, 1):
    wk.column_dimensions[L(i)].width = b
    c = wk.cell(HR, i, k); c.font = font(bold=True, color=ZWART); c.fill = fill(LICHT)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    wk.cell(GR, i).fill = fill(ZWART)
wk.merge_cells(start_row=GR, start_column=1, end_row=GR, end_column=len(KK))
c = wk.cell(GR, 1, "KOSTEN  ·  wat Q4S zelf betaalt"); c.font = font(bold=True, color="FFFFFF", size=9)
c.alignment = Alignment(horizontal="center")
wk.row_dimensions[HR].height = 30
kosten.sort(key=lambda x: datum(x["ont"])[0] or dt.datetime(2100, 1, 1))
eind_k = D0 + len(kosten) + (800 if SJABLOON else EXTRA) - 1
for n in range(D0, eind_k + 1):
    x = kosten[n - D0] if n - D0 < len(kosten) else None
    btwv = "f"
    if x:
        d, o1 = datum(x["ont"]); b, o2 = datum(x["bet"])
        oms = "; ".join(s for s in [x["klant"], x["opm"], o1 and f"datum: {o1}", o2 and f"betaald: {o2}"] if s)
        for col, v in {1: d, 3: x["naam"], 4: oms or None, 5: soort(x), 6: x["fnr"], 7: x["ex_in"],
                       8: x["pct_in"], 11: b}.items():
            if v is not None: wk.cell(n, col, v)
        btwv = btw_bedrag(x["btw_in"], x["ex_in"], x["pct_in"], None)
        if btwv != "f": wk.cell(n, 9, btwv)
    f = lambda s: s.replace("#", str(n))
    wk.cell(n, 2, f('=IF(A#="","",YEAR(A#)&" Q"&ROUNDUP(MONTH(A#)/3,0))'))
    if btwv == "f": wk.cell(n, 9, f('=IF(G#="","",ROUND(G#*H#,2))'))
    wk.cell(n, 10, f('=IF(G#="","",G#+N(I#))'))
    wk.cell(n, 12, f('=IF(G#="","",IF(OR(K#="",K#>TODAY()),"Nog betalen","Betaald"))'))
    for col in range(1, len(KK) + 1):
        c = wk.cell(n, col); c.font = font(color=ZWART); c.border = rand
        if col in (1, 11): c.number_format = DATUM
        if col in (7, 9, 10): c.number_format = EURO
        if col == 8: c.number_format = PCT
        if col in (2, 9, 10, 12): c.font = font(color=INKT); c.fill = fill("FAFAFA")
tab = Table(displayName="tblKosten", ref=f"A{HR}:{L(len(KK))}{eind_k}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
wk.add_table(tab)
wk.freeze_panes = f"D{D0}"
for s, (bg, fg) in {"Betaald": STATUS["Afgerond"], "Nog betalen": STATUS["Nog betalen"]}.items():
    wk.conditional_formatting.add(f"L{D0}:L{eind_k}", CellIsRule(
        operator="equal", formula=[f'"{s}"'], fill=fill(bg), font=Font(name=F, size=10, bold=True, color=fg)))
dv = DataValidation(type="list", formula1=f"Lijsten!$A$2:$A${len(SOORTEN)+1}", allow_blank=True)
wk.add_data_validation(dv); dv.add(f"E{D0}:E{eind_k}")
pct = PCTV(); wk.add_data_validation(pct); pct.add(f"H{D0}:H{eind_k}")
dat = DATV(); wk.add_data_validation(dat); dat.add(f"A{D0}:A{eind_k}"); dat.add(f"K{D0}:K{eind_k}")

# ============================================================================
# OVERZICHT — het dashboard (staat vooraan)
# ============================================================================
wo2 = wb.create_sheet("Overzicht", 0)
for i, b in enumerate([3, 16, 18, 18, 18, 18, 18, 17, 3], 1): wo2.column_dimensions[L(i)].width = b
wo2.sheet_view.showGridLines = False
for r in range(1, 5): wo2.row_dimensions[r].height = 20
logo(wo2, 78)
c = wo2.cell(1, 3, "Factuuroverzicht"); c.font = font(size=22, bold=True, color=ZWART)
c = wo2.cell(2, 3, "Q4S B.V.  ·  bedragen ex btw, tenzij anders vermeld"); c.font = font(size=10, color=GRIJS)
c = wo2.cell(4, 7, "Jaar"); c.font = font(bold=True, color=GRIJS); c.alignment = Alignment(horizontal="right", vertical="center")
J = wo2.cell(4, 8, 2026); J.font = font(size=14, bold=True, color="FFFFFF"); J.fill = fill(ZWART)
J.alignment = Alignment(horizontal="center", vertical="center")
jv = DataValidation(type="list", formula1='"2025,2026,2027,2028"'); wo2.add_data_validation(jv); jv.add("H4")
wo2.cell(5, 8, "▲ kies hier het jaar").font = font(size=8, color=GRIJS, italic=True)
wo2["H5"].alignment = Alignment(horizontal="center")
for col in range(2, 9): wo2.cell(6, col).border = Border(bottom=Side(style="medium", color=ZWART))
JAAR = "$H$4"

def tussen(y, m1, m2):  # SUMIFS-criteria op een datumkolom
    return f'">="&DATE({y},{m1},1),"<"&DATE({y},{m2},1)'
def som(blad, waarde, datumkol, m1, m2):
    rng = f"{blad}!${datumkol}${D0}:${datumkol}${MAXR}"
    return (f"SUMIFS({blad}!${waarde}${D0}:${waarde}${MAXR},{rng},\">=\"&DATE({JAAR},{m1},1),"
            f"{rng},\"<\"&DATE({JAAR},{m2},1))")

# KPI-tegels
TEGELS = [("Omzet", "D"), ("Inkoop freelancers", "E"), ("Brutomarge", "F"), ("Kosten", "G"), ("Winst", "H")]
# tegels: 5 tegels op rijen 8-10, elk één kolom breed (B..H met tussenruimte → we gebruiken B,C,D,E,G? houd het simpel)
tegelkol = [3, 4, 5, 6, 7]   # elke tegel precies boven zijn kolom in de kwartaaltabel
QR0 = 15  # eerste kwartaalrij
totrij = QR0 + 4
for (label, bron), col in zip(TEGELS, tegelkol):
    kolletter = L(col)
    a = wo2.cell(8, col, label.upper()); a.font = font(size=8, bold=True, color=GRIJS)
    v = wo2.cell(9, col, f"={kolletter}{totrij}"); v.number_format = '€ #,##0;[Red]-€ #,##0;"€ 0"'
    v.font = font(size=16, bold=True, color="16A34A" if label == "Winst" else ZWART)
    for r in (8, 9, 10):
        wo2.cell(r, col).fill = fill(LICHT)
        wo2.cell(r, col).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        wo2.cell(r, col).border = Border(left=Side(style="thick", color="FFFFFF"), right=Side(style="thick", color="FFFFFF"))
    wo2.cell(10, col).border = Border(left=Side(style="thick", color="FFFFFF"), right=Side(style="thick", color="FFFFFF"),
                                      bottom=Side(style="thick", color="16A34A" if label == "Winst" else ZWART))
wo2.row_dimensions[9].height = 30
wo2.cell(10, 2).value = None

# Kwartaaltabel
c = wo2.cell(12, 2, "Per kwartaal"); c.font = font(size=13, bold=True, color=ZWART)
KOP = ["Kwartaal", "Omzet", "Inkoop freelancers", "Brutomarge", "Kosten", "Winst", "Btw afdragen"]
for i, k in enumerate(KOP, 2):
    c = wo2.cell(14, i, k or None); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(ZWART if k else "FFFFFF")
    c.alignment = Alignment(horizontal="left" if i == 2 else "right", vertical="center")
wo2.row_dimensions[14].height = 22
for q in range(1, 5):
    r = QR0 + q - 1; m1, m2 = 3 * q - 2, 3 * q + 1
    m2s = f"{m2}" if m2 <= 12 else "13"   # DATE(j,13,1) = 1 jan volgend jaar
    wo2.cell(r, 2, f"Q{q}")
    wo2.cell(r, 3, "=" + som("Facturen", "M", "Q", m1, m2s))
    wo2.cell(r, 4, "=" + som("Facturen", "G", "A", m1, m2s))
    wo2.cell(r, 5, f"=C{r}-D{r}")
    wo2.cell(r, 6, "=" + som("Kosten", "G", "A", m1, m2s))
    wo2.cell(r, 7, f"=E{r}-F{r}")
    wo2.cell(r, 8, "=" + som("Facturen", "O", "Q", m1, m2s) + "-" + som("Facturen", "I", "A", m1, m2s)
             + "-" + som("Kosten", "I", "A", m1, m2s))
r = totrij
wo2.cell(r, 2, "Totaal")
for col in (3, 4, 5, 6, 7, 8): wo2.cell(r, col, f"=SUM({L(col)}{QR0}:{L(col)}{QR0+3})")
for rr in range(QR0, totrij + 1):
    for col in range(2, 9):
        c = wo2.cell(rr, col); c.border = Border(bottom=dun)
        c.font = font(bold=(rr == totrij), color=ZWART, size=11)
        if col > 2: c.number_format = EURO; c.alignment = Alignment(horizontal="right")
        if rr == totrij: c.fill = fill(LICHT); c.border = Border(top=Side(style="medium", color=ZWART))
    wo2.row_dimensions[rr].height = 20
wo2.conditional_formatting.add(f"G{QR0}:G{totrij}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, color="16A34A", bold=True)))

# Openstaand
c = wo2.cell(22, 2, "Wat staat er nog open?"); c.font = font(size=13, bold=True, color=ZWART)
c = wo2.cell(23, 2, "incl btw, alle jaren"); c.font = font(size=9, color=GRIJS)
OPEN = [
    ("Klanten moeten ons nog betalen", 'SUMIFS(Facturen!$P$6:$P$2000,Facturen!$T$6:$T$2000,"Wacht op klant")',
     'COUNTIF(Facturen!$T$6:$T$2000,"Wacht op klant")', "DBEAFE"),
    ("Wij moeten freelancers nog betalen", 'SUMIFS(Facturen!$J$6:$J$2000,Facturen!$T$6:$T$2000,"Nog betalen")',
     'COUNTIF(Facturen!$T$6:$T$2000,"Nog betalen")', "FEF3C7"),
    ("Wij moeten kosten nog betalen", 'SUMIFS(Kosten!$J$6:$J$2000,Kosten!$L$6:$L$2000,"Nog betalen")',
     'COUNTIF(Kosten!$L$6:$L$2000,"Nog betalen")', "FEF3C7"),
    ("Freelancerfacturen nog niet doorgefactureerd", 'SUMIFS(Facturen!$J$6:$J$2000,Facturen!$T$6:$T$2000,"Nog factureren")',
     'COUNTIF(Facturen!$T$6:$T$2000,"Nog factureren")', "FEE2E2"),
]
for i, (label, bedrag, aantal, kleur) in enumerate(OPEN):
    r = 24 + i
    wo2.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    a = wo2.cell(r, 2, label); a.font = font(size=11, color=ZWART); a.fill = fill(kleur)
    a.alignment = Alignment(vertical="center", indent=1)
    n = wo2.cell(r, 7, f'={aantal}&IF({aantal}=1," factuur"," facturen")'); n.font = font(color=GRIJS)
    n.alignment = Alignment(horizontal="right", vertical="center")
    b = wo2.cell(r, 8, f"={bedrag}"); b.number_format = EURO; b.font = font(size=11, bold=True, color=ZWART)
    b.alignment = Alignment(vertical="center")
    for col in range(2, 9): wo2.cell(r, col).border = Border(bottom=Side(style="thin", color="FFFFFF"))
    for col in (6, 7, 8): wo2.cell(r, col).fill = fill(kleur)
    wo2.row_dimensions[r].height = 22

# Uitleg
c = wo2.cell(30, 2, "Zo werkt het"); c.font = font(size=13, bold=True, color=ZWART)
UITLEG = [
    "1.  Factuur van een freelancer binnen?  →  tabblad Facturen, nieuwe regel: datum, naam, klant, factuurnr en bedrag ex btw.",
    "2.  Onze factuur aan de klant gemaakt?  →  zelfde regel: Q4S factuurnr, bedrag ex btw en de datum verstuurd.",
    "     Eén Q4S-factuur voor meerdere freelancerfacturen? Zet hetzelfde Q4S-nummer op elke regel en het bedrag alleen op de eerste.",
    "3.  Betaald of geld binnen?  →  vul 'Betaald op' / 'Klant betaald op' in. De Status springt vanzelf naar Afgerond.",
    "4.  Eigen kosten (huur, software, tankpas …)  →  tabblad Kosten, kies een Soort.",
    "Grijze vakjes rekenen zelf — niet overschrijven.  Omzet telt op 'Verstuurd op', inkoop en kosten op de factuurdatum.",
]
for i, t in enumerate(UITLEG):
    c = wo2.cell(31 + i, 2, t); c.font = font(size=10, color=INKT if i < 5 else GRIJS, italic=(i == 5))

# Lijsten
wl = wb.create_sheet("Lijsten")
wl["A1"] = "Soorten kosten"; wl["A1"].font = font(bold=True)
for i, s in enumerate(SOORTEN, 2): wl.cell(i, 1, s).font = font()
wl.column_dimensions["A"].width = 32
wl.sheet_state = "hidden"

for ws in (wo2, wf, wk):
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1 if ws is wo2 else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if ws is not wo2: ws.print_title_rows = f"{GR}:{HR}"
wo2.sheet_properties.tabColor = ZWART; wf.sheet_properties.tabColor = "3F3F46"; wk.sheet_properties.tabColor = "6B6B70"
wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(NIEUW)
print(f"facturen {len(facturen)}  kosten {len(kosten)}  freelancers {sorted(freelancers)}")
